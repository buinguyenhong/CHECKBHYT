import sys
import os
import datetime

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "web_app"))

from xml_validator_v2.source_adapter import XmlSourceAdapter
from xml_validator_v2.validation_service import run_validation_job
from xml_validator_v2.db import ValSessionLocal, init_val_db
from xml_validator_v2.models import ValRun, ValResult, ValidationError
from xml_validator_v2.report_exporter import export_run_to_excel

class MockXmlSourceAdapter(XmlSourceAdapter):
    """Mock Adapter trả về dữ liệu XML giả lập cho unit test"""
    def fetch_xml_batch(self, ma_lk_list, chunk_size=50):
        res = {k: {} for k in ma_lk_list}
        # Hồ sơ 1: Hợp lệ 9324
        if "MOCK_LK_VALID_9324" in res:
            res["MOCK_LK_VALID_9324"] = {
                "XML1": [{
                    "MA_LK": "MOCK_LK_VALID_9324",
                    "MA_LYDO_VVIEN": "1",
                    "MA_CSKCB": "66232",
                    "MA_DKBD": "66232",
                    "MA_THE": "DN4662320123456",
                    "T_TONGCHI": 300000,
                    "T_BNTT": 50000,
                    "T_BHTT": 250000,
                    "T_NGUONKHAC": 0
                }],
                "XML2": [{"SO_LUONG": 2, "DON_GIA": 50000, "TYLE_TT": 100, "THANH_TIEN": 100000}],
                "XML3": [{"SO_LUONG": 1, "DON_GIA": 200000, "TYLE_TT": 100, "THANH_TIEN": 200000, "TEN_DICH_VU": "XÉT NGHIỆM"}]
            }
        # Hồ sơ 2: Có lỗi vi phạm chuẩn 4750 (Khám bệnh có tiền giường + T_BHTT = 0)
        if "MOCK_LK_ERR_4750" in res:
            res["MOCK_LK_ERR_4750"] = {
                "XML1": [{
                    "MA_LK": "MOCK_LK_ERR_4750",
                    "T_TONGCHI_BV": 150000,
                    "T_BNTT": 150000,
                    "T_BNCCT": 0,
                    "T_BHTT": 0,  # Vi phạm: T_BHTT = 0
                    "T_NGUONKHAC": 0,
                    "MA_LOAI_KCB": "1",  # Khám bệnh
                    "GIOI_TINH": "1",
                    "NGAY_VAO": "202610200800",
                    "NGAY_RA": "202610201100",
                    "NGAY_SINH": "19900510",
                    "MA_THE": "DN4662320123456"
                }],
                "XML3": [
                    {"MA_NHOM": "14", "SO_LUONG": 1, "DON_GIA": 150000, "TYLE_TT": 100, "THANH_TIEN": 150000, "T_BHTT": 0}
                ]
            }
        case_table_status = {k: {f"XML{i}": "SUCCESS" for i in range(1, 16)} for k in ma_lk_list}
        batch_table_status = {f"XML{i}": {"status": "SUCCESS", "rows": 10} for i in range(1, 16)}
        return res, case_table_status, batch_table_status

def test_full_validation_run():
    init_val_db()
    mock_adapter = MockXmlSourceAdapter()

    cases_input = [
        {"ma_lk": "MOCK_LK_VALID_9324", "ho_ten": "Nguyễn Văn A", "loai_ca": "Ngoại trú", "ma_the": "DN4662320123456", "ngay_ra_vien": datetime.date(2026, 10, 20)},
        {"ma_lk": "MOCK_LK_ERR_4750", "ho_ten": "Trần Thị B", "loai_ca": "Ngoại trú", "ma_the": "DN4662320123456", "ngay_ra_vien": datetime.date(2026, 10, 20)},
        {"ma_lk": "MOCK_LK_MISSING", "ho_ten": "Lê Văn C", "loai_ca": "Nội trú", "ma_the": "TE1662320123456", "ngay_ra_vien": datetime.date(2026, 10, 21)}
    ]

    # Chạy job
    run_validation_job(
        cases_meta=cases_input,
        scope_type="UNSENT_BATCH",
        username="test_runner",
        xml_adapter=mock_adapter
    )

    # Kiểm tra trong DB SQLite riêng
    db = ValSessionLocal()
    try:
        latest_run = db.query(ValRun).order_by(ValRun.id.desc()).first()
        assert latest_run is not None
        assert latest_run.status == "COMPLETED"
        assert latest_run.total_cases == 3
        assert latest_run.processed_cases == 3
        assert latest_run.has_error_cases == 1
        assert latest_run.incomplete_cases == 1  # Không bị gán nhầm NO_ERRORS_FOUND vì chuẩn 9324 còn quy tắc chưa có code
        assert latest_run.missing_xml_cases == 1

        results = db.query(ValResult).filter(ValResult.run_id == latest_run.id).all()
        assert len(results) == 3

        res_map = {r.ma_lk: r for r in results}
        assert res_map["MOCK_LK_VALID_9324"].overall_status == "INCOMPLETE_DATA"
        assert res_map["MOCK_LK_VALID_9324"].violated_rules_count == 0
        assert res_map["MOCK_LK_VALID_9324"].incomplete_rules_count > 0

        assert res_map["MOCK_LK_ERR_4750"].overall_status == "HAS_ERRORS"
        assert res_map["MOCK_LK_ERR_4750"].violated_rules_count >= 1

        assert res_map["MOCK_LK_MISSING"].overall_status == "MISSING_XML"

        # Kiểm tra xuất Excel: Xác minh Sheet 1 xuất đủ 3 dòng kết quả (chứng minh bug ValResult.id đã được sửa triệt để)
        out_excel = os.path.join(os.path.dirname(__file__), "..", "web_app", "exports", "test_report.xlsx")
        export_run_to_excel(db, latest_run.id, out_excel)
        assert os.path.exists(out_excel)
        assert os.path.getsize(out_excel) > 0

        import openpyxl
        wb = openpyxl.load_workbook(out_excel)
        ws1 = wb["Tổng Hợp Hồ Sơ"]
        # Có 1 dòng header + 3 dòng kết quả = 4 dòng
        assert ws1.max_row == 4, f"Expected 4 rows in Excel Sheet 1 but got {ws1.max_row}"
        print(f"Export Excel test PASSED! File size: {os.path.getsize(out_excel)} bytes, Max row: {ws1.max_row}")

    finally:
        db.close()

if __name__ == "__main__":
    test_full_validation_run()
    print("ALL INTEGRATION TESTS PASSED SUCCESSFULLY!")
