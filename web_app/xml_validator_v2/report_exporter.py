import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from .models import ValRun, ValResult, ValidationError

def export_run_to_excel(db_session, run_id: int, output_filepath: str) -> str:
    """
    Xuất báo cáo kết quả kiểm tra XML ra file Excel định dạng chuẩn.
    Sheet 1: Tổng hợp hồ sơ
    Sheet 2: Chi tiết các quy tắc vi phạm
    """
    run = db_session.query(ValRun).filter(ValRun.id == run_id).first()
    if not run:
        raise ValueError(f"Không tìm thấy đợt kiểm tra ID: {run_id}")

    results = db_session.query(ValResult).filter(ValResult.run_id == run_id).all()
    errors = db_session.query(ValidationError).filter(ValidationError.run_id == run_id).all()

    wb = openpyxl.Workbook()
    # Sheet 1: Tong hop
    ws1 = wb.active
    ws1.title = "Tổng Hợp Hồ Sơ"

    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    border_thin = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )

    headers_ws1 = [
        "STT", "Mã Liên Kết", "Họ Tên", "Loại Ca", "Mã Thẻ", "Ngày Ra Viện",
        "Chuẩn Áp Dụng", "Trạng Thái Validate", "Số Lỗi", "Số Đạt", "Chưa Đánh Giá", "Tóm Tắt Kết Quả"
    ]
    ws1.append(headers_ws1)
    for col_idx in range(1, len(headers_ws1) + 1):
        cell = ws1.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center

    for idx, r in enumerate(results, 1):
        row_data = [
            idx, r.ma_lk, r.ho_ten, r.loai_ca, r.ma_the,
            r.ngay_ra_vien.strftime("%d/%m/%Y") if r.ngay_ra_vien else "",
            r.applied_standard, r.overall_status,
            r.violated_rules_count, r.passed_rules_count, r.incomplete_rules_count,
            r.summary_text
        ]
        ws1.append(row_data)
        for col_idx in range(1, len(row_data) + 1):
            cell = ws1.cell(row=idx + 1, column=col_idx)
            cell.border = border_thin
            if col_idx in [1, 4, 6, 7, 8, 9, 10, 11]:
                cell.alignment = align_center
            else:
                cell.alignment = align_left

    # Sheet 2: Chi tiet loi
    ws2 = wb.create_sheet(title="Chi Tiết Vi Phạm")
    headers_ws2 = [
        "STT", "Mã Liên Kết", "Chuẩn", "Mã Lỗi", "Tên Quy Tắc",
        "Bảng XML", "Trường Dữ Liệu", "Giá Trị Thực Tế", "Điều Kiện Kỳ Vọng", "Nội Dung Diễn Giải"
    ]
    ws2.append(headers_ws2)
    for col_idx in range(1, len(headers_ws2) + 1):
        cell = ws2.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center

    for idx, err in enumerate(errors, 1):
        row_data = [
            idx, err.ma_lk, err.standard, err.rule_code, err.rule_name,
            err.xml_table, err.field_name, err.actual_value, err.expected_condition, err.message
        ]
        ws2.append(row_data)
        for col_idx in range(1, len(row_data) + 1):
            cell = ws2.cell(row=idx + 1, column=col_idx)
            cell.border = border_thin
            if col_idx in [1, 3, 4, 6, 7]:
                cell.alignment = align_center
            else:
                cell.alignment = align_left

    # Tự động căn chỉnh độ rộng cột
    for ws in [ws1, ws2]:
        for col in ws.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                val = str(cell.value or '')
                max_len = max(max_len, len(val))
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

    os.makedirs(os.path.dirname(output_filepath), exist_ok=True)
    wb.save(output_filepath)
    return output_filepath
