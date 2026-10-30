import re
from decimal import Decimal, InvalidOperation
import openpyxl
import pandas as pd
from io import BytesIO
from typing import Union, Optional, Set, Dict, Iterable
from services.his_service import chuan_hoa_ma_lk, parse_datetime_to_date

MA_LK_CANDIDATE_HEADERS = {
    "ma_lk", "mã liên kết", "mã_lk", "ma lk", "so_seri"
}


def extract_cell_text_openpyxl(cell) -> str:
    """
    Đọc trực tiếp giá trị từ đối tượng Cell của openpyxl, kiểm tra:
    - cell.value
    - cell.data_type
    - cell.number_format
    Nguyên tắc:
    1. Nếu ô là chuỗi (data_type in ('s', 'str', 'inlineStr') hoặc isinstance(val, str)):
       Giữ nguyên chuỗi gốc (bảo toàn số 0 ở cuối như '26.035530').
    2. Nếu ô là kiểu số (data_type == 'n' hoặc int/float):
       - Nếu cell.number_format có khai báo phần thập phân cố định (ví dụ '0.000000', '#,##0.000000'):
         Khôi phục đúng số chữ số thập phân theo định dạng hiển thị của ô đó.
       - Nếu cell.number_format là 'General' hoặc '@':
         Tuyệt đối KHÔNG tự động thêm số 0 nếu không có căn cứ định dạng.
    """
    val = cell.value
    if val is None:
        return ""
    if cell.data_type in ("s", "str", "inlineStr") or isinstance(val, str):
        return str(val).strip()

    if cell.data_type == "n" or isinstance(val, (int, float, Decimal)):
        fmt = str(getattr(cell, "number_format", "") or "").strip()
        # Lấy phần định dạng cho số dương (trước dấu ';')
        pos_fmt = fmt.split(";")[0].strip() if fmt else ""
        if pos_fmt and pos_fmt.lower() != "general" and pos_fmt != "@":
            m = re.search(r"[0#]\.([0#]+)", pos_fmt)
            if m:
                dec_mask = m.group(1)
                min_decimals = len(dec_mask.rstrip("#"))
                if min_decimals > 0:
                    try:
                        d_val = Decimal(str(val))
                        quant = Decimal("1." + ("0" * min_decimals))
                        return str(d_val.quantize(quant))
                    except (InvalidOperation, ValueError):
                        pass
        if isinstance(val, float) and val.is_integer():
            return str(int(val))
        return str(val).strip()

    if isinstance(val, float) and pd.isna(val):
        return ""
    return str(val).strip()


def read_excel_preserving_ma_lk(file_source: Union[str, BytesIO, bytes]) -> pd.DataFrame:
    """
    Đọc tệp Excel sang DataFrame với dtype=str, đồng thời dùng openpyxl kiểm tra trực tiếp
    cell.value, cell.data_type, cell.number_format trên cột MA_LK / Mã liên kết để khôi phục
    chính xác các ô số có định dạng hiển thị thập phân (ví dụ number_format='0.000000').
    """
    if isinstance(file_source, bytes):
        stream_for_pd = BytesIO(file_source)
        stream_for_wb = BytesIO(file_source)
    elif isinstance(file_source, BytesIO):
        raw_bytes = file_source.getvalue()
        stream_for_pd = BytesIO(raw_bytes)
        stream_for_wb = BytesIO(raw_bytes)
    else:
        stream_for_pd = file_source
        stream_for_wb = file_source

    df = pd.read_excel(stream_for_pd, dtype=str)
    if df.empty:
        return df

    try:
        wb = openpyxl.load_workbook(stream_for_wb, data_only=True)
        ws = wb.active
        header_cells = next(ws.iter_rows(min_row=1, max_row=1, values_only=False), [])
        lk_col_positions = []
        for idx, h_cell in enumerate(header_cells):
            h_name = str(h_cell.value or "").strip()
            if h_name.lower() in MA_LK_CANDIDATE_HEADERS and h_name in df.columns:
                lk_col_positions.append((idx + 1, h_name))

        if lk_col_positions:
            for col_1based, col_name in lk_col_positions:
                extracted_vals = []
                for row in ws.iter_rows(min_row=2, min_col=col_1based, max_col=col_1based, values_only=False):
                    extracted_vals.append(extract_cell_text_openpyxl(row[0]))
                if len(extracted_vals) == len(df):
                    df[col_name] = extracted_vals
                elif len(extracted_vals) > len(df):
                    df[col_name] = extracted_vals[:len(df)]
        wb.close()
    except Exception:
        pass

    return df


def build_trusted_numeric_ma_lk_map(trusted_keys: Iterable[str]) -> Dict[str, Optional[str]]:
    """
    Xây dựng bảng ánh xạ từ dạng chuẩn hóa số thập phân (đã bỏ số 0 cuối) sang mã MA_LK gốc
    từ nguồn dữ liệu đáng tin cậy (SQL Server HIS / bảng Record).
    Nếu có từ 2 mã khác nhau cùng rút gọn về 1 giá trị (nhập nhằng), đánh dấu None để không tự đoán.
    """
    numeric_map: Dict[str, Optional[str]] = {}
    for raw_k in trusted_keys:
        if not raw_k:
            continue
        k = chuan_hoa_ma_lk(raw_k).upper()
        if re.match(r"^\d+\.\d+$", k):
            canon = k.rstrip("0").rstrip(".")
            if canon not in numeric_map:
                numeric_map[canon] = k
            elif numeric_map[canon] != k:
                numeric_map[canon] = None
    return numeric_map


def resolve_ma_lk_against_trusted_set(
    ma_lk: str,
    trusted_keys: Set[str],
    numeric_map: Optional[Dict[str, Optional[str]]] = None
) -> str:
    """
    Đối chiếu mã MA_LK đọc từ Excel với tập mã nguồn đáng tin cậy (SQL Server HIS):
    - Nếu khớp trực tiếp trong trusted_keys -> giữ nguyên.
    - Nếu là mã dạng số thập phân bị mất số 0 ở cuối (do Excel lưu kiểu số General) và trong
      trusted_keys có DUY NHẤT 1 mã khớp chính xác giá trị số đó -> khôi phục theo mã chuẩn của HIS.
    - Nếu không có căn cứ trong trusted_keys hoặc bị trùng/nhập nhằng -> giữ nguyên, tuyệt đối không tự thêm số 0.
    """
    lk = chuan_hoa_ma_lk(ma_lk).upper()
    if not lk or not trusted_keys:
        return lk
    if lk in trusted_keys:
        return lk
    if re.match(r"^\d+\.\d+$", lk):
        if numeric_map is None:
            numeric_map = build_trusted_numeric_ma_lk_map(trusted_keys)
        canon = lk.rstrip("0").rstrip(".")
        candidate = numeric_map.get(canon)
        if candidate and candidate.startswith(lk):
            return candidate
    return lk


def load_listbh(
    file_source: Union[str, BytesIO],
    key_col: str = "Mã liên kết",
    date_col: str = "Ngày ra",
    trusted_keys: Optional[Set[str]] = None
) -> pd.DataFrame:
    """
    Đọc file listbh.xlsx từ file path hoặc luồng byte, chuẩn hoá cột và bảo toàn MA_LK
    """
    df = read_excel_preserving_ma_lk(file_source)
    
    found_key = None
    candidates_key = [key_col, "Mã liên kết", "MA_LK", "MÃ LIÊN KẾT", "Mã Liên Kết", "MÃ_LK", "Ma_LK", "Mã LK", "SO_SERI"]
    for c in candidates_key:
        if c in df.columns:
            found_key = c
            break
    if not found_key:
        col_map = {str(c).lower(): c for c in df.columns}
        for c in ["mã liên kết", "ma_lk", "mã_lk", "ma lk"]:
            if c in col_map:
                found_key = col_map[c]
                break
    if not found_key:
        raise ValueError(f"Tệp danh sách BHYT thiếu cột khóa chính 'Mã liên kết' / 'MA_LK'. Các cột hiện có: {list(df.columns)}")
        
    df[found_key] = df[found_key].apply(chuan_hoa_ma_lk)
    if trusted_keys:
        num_map = build_trusted_numeric_ma_lk_map(trusted_keys)
        df[found_key] = df[found_key].apply(lambda x: resolve_ma_lk_against_trusted_set(x, trusted_keys, num_map))

    found_date = None
    candidates_date = [date_col, "Ngày ra", "NGÀY RA", "NgayRa", "NGAY_RA", "Ngày ra viện", "NGÀY RA VIỆN", "Ngay_Ra"]
    for c in candidates_date:
        if c in df.columns:
            found_date = c
            break
    if not found_date:
        col_map = {str(c).lower(): c for c in df.columns}
        for c in ["ngày ra", "ngayra", "ngay_ra", "ngày ra viện"]:
            if c in col_map:
                found_date = col_map[c]
                break

    if found_date:
        df["_ngay"] = parse_datetime_to_date(df[found_date])
    else:
        df["_ngay"] = pd.NaT

    out = df[[found_key, "_ngay"]].copy()
    out = out.rename(columns={found_key: "MA_LK"})
    return out

def load_hosoloichitiet(
    file_source: Union[str, BytesIO],
    trusted_keys: Optional[Set[str]] = None
) -> pd.DataFrame:
    """
    Đọc file lỗi HoSoLoiChiTiet.xlsx từ file path hoặc luồng byte, kiểm tra trực tiếp openpyxl
    """
    df = read_excel_preserving_ma_lk(file_source)
    found_key = None
    for c in ["MA_LK", "Mã liên kết", "MÃ LIÊN KẾT", "MÃ_LK", "Ma_LK"]:
        if c in df.columns:
            found_key = c
            break
    if not found_key:
        col_map = {str(c).lower(): c for c in df.columns}
        for c in ["ma_lk", "mã liên kết", "mã_lk"]:
            if c in col_map:
                found_key = col_map[c]
                break
    if not found_key:
        raise ValueError(f"Tệp lỗi HoSoLoiChiTiet.xlsx phải chứa cột bắt buộc 'MA_LK'. Các cột hiện có: {list(df.columns)}")

    if found_key != "MA_LK":
        df = df.rename(columns={found_key: "MA_LK"})

    found_loi = None
    for c in ["MALOI", "Mã lỗi", "MÃ LỖI", "MaLoi"]:
        if c in df.columns:
            found_loi = c
            break
    if found_loi and found_loi != "MALOI":
        df = df.rename(columns={found_loi: "MALOI"})
    elif "MALOI" not in df.columns:
        df["MALOI"] = ""

    found_mota = None
    for c in ["MOTALOI", "Mô tả lỗi", "MÔ TẢ LỖI", "MoTaLoi", "Nội dung lỗi"]:
        if c in df.columns:
            found_mota = c
            break
    if found_mota and found_mota != "MOTALOI":
        df = df.rename(columns={found_mota: "MOTALOI"})
    elif "MOTALOI" not in df.columns:
        df["MOTALOI"] = ""

    df["MALOI"] = df["MALOI"].fillna("")
    df["MOTALOI"] = df["MOTALOI"].fillna("")
    df["MA_LK"] = df["MA_LK"].apply(chuan_hoa_ma_lk)
    if trusted_keys:
        num_map = build_trusted_numeric_ma_lk_map(trusted_keys)
        df["MA_LK"] = df["MA_LK"].apply(lambda x: resolve_ma_lk_against_trusted_set(x, trusted_keys, num_map))

    found_date = None
    for c in ["Ngày ra", "NGÀY RA", "NgayRa", "NGAY_RA", "Ngày ra viện"]:
        if c in df.columns:
            found_date = c
            break
    if found_date:
        df["Ngày ra"] = parse_datetime_to_date(df[found_date])

    out_cols = ["MA_LK", "MALOI", "MOTALOI"]
    if "Ngày ra" in df.columns:
        out_cols.append("Ngày ra")

    return df[out_cols].copy()

def filter_listbh_by_date(df: pd.DataFrame, tu_ngay, den_ngay) -> pd.DataFrame:
    """
    Lọc danh sách đã gửi BHYT trong khoảng ngày được chọn
    """
    if df.empty:
        return df
    if "_ngay" not in df.columns:
        return df
    # Đảm bảo tu_ngay và den_ngay là đối tượng date
    mask = (df["_ngay"] >= tu_ngay) & (df["_ngay"] <= den_ngay)
    return df.loc[mask].copy()
